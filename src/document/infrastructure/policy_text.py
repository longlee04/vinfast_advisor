"""Bounded text extraction for synchronous policy analysis."""

from __future__ import annotations

from html.parser import HTMLParser
from io import BytesIO
from zipfile import BadZipFile

import anyio
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from src.document.application.errors import DocumentTextUnavailableError, UnsupportedPolicyDocumentError

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_MIME = "text/plain"
HTML_MIME = "text/html"


class _PolicyHtmlTextExtractor(HTMLParser):
    """Keep policy section structure while dropping navigation and executable content."""

    _SKIPPED_TAGS = frozenset({"script", "style", "nav", "noscript", "svg"})
    _BLOCK_TAGS = frozenset({"p", "div", "li", "section", "article", "br"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._heading_level: int | None = None
        self._anchor_href: str | None = None
        self._cell_parts: list[str] | None = None
        self._row_cells: list[str] | None = None
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        normalized = tag.casefold()
        if normalized in self._SKIPPED_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if normalized in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._heading_level = int(normalized[1])
            self.parts.append("\n")
        elif normalized == "tr":
            self._row_cells = []
        elif normalized in {"th", "td"} and self._row_cells is not None:
            self._cell_parts = []
        elif normalized == "a":
            self._anchor_href = next((value for key, value in attrs if key == "href"), None)
        elif normalized in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.casefold()
        if normalized in self._SKIPPED_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if normalized in {"th", "td"} and self._cell_parts is not None:
            if self._row_cells is not None:
                self._row_cells.append(" ".join(self._cell_parts).strip())
            self._cell_parts = None
        elif normalized == "tr" and self._row_cells is not None:
            row = " | ".join(cell for cell in self._row_cells if cell)
            if row:
                self.parts.extend((row, "\n"))
            self._row_cells = None
        elif normalized == "a":
            if self._anchor_href:
                self.parts.append(f" ({self._anchor_href})")
            self._anchor_href = None
        elif normalized in self._BLOCK_TAGS or normalized in {
            "h1", "h2", "h3", "h4", "h5", "h6"
        }:
            self.parts.append("\n")
            self._heading_level = None

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        value = " ".join(data.split())
        if not value:
            return
        if self._cell_parts is not None:
            self._cell_parts.append(value)
            return
        if self._heading_level is not None:
            self.parts.append(f"{'#' * self._heading_level} {value}")
            self._heading_level = None
        else:
            self.parts.append(value)


class PolicyDocumentTextExtractor:
    """Extract meaningful text from PDF, DOCX, HTML, or UTF-8 TXT without OCR."""

    async def extract(self, payload: bytes, *, filename: str, content_type: str) -> str:
        """Parse one stored object off the event loop and normalize whitespace."""
        try:
            text = await anyio.to_thread.run_sync(
                self._extract_sync,
                payload,
                filename,
                content_type,
            )
        except (PdfReadError, UnicodeDecodeError, BadZipFile, PackageNotFoundError, ValueError) as error:
            raise DocumentTextUnavailableError() from error
        normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
        if not normalized:
            raise DocumentTextUnavailableError()
        return normalized

    def _extract_sync(self, payload: bytes, filename: str, content_type: str) -> str:
        lower_name = filename.lower()
        if content_type == PDF_MIME and lower_name.endswith(".pdf"):
            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages)
        if content_type == DOCX_MIME and lower_name.endswith(".docx"):
            document = DocxDocument(BytesIO(payload))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_rows = [
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            return "\n".join([*paragraphs, *table_rows])
        if content_type == TEXT_MIME and lower_name.endswith(".txt"):
            return payload.decode("utf-8")
        if content_type == HTML_MIME:
            parser = _PolicyHtmlTextExtractor()
            parser.feed(payload.decode("utf-8"))
            parser.close()
            return "".join(parser.parts)
        raise UnsupportedPolicyDocumentError()
