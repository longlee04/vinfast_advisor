"""Text extraction boundaries for policy analysis."""

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.document.application.errors import DocumentTextUnavailableError, UnsupportedPolicyDocumentError
from src.document.infrastructure.policy_text import PolicyDocumentTextExtractor


@pytest.mark.asyncio
async def test_extracts_utf8_plain_text() -> None:
    text = await PolicyDocumentTextExtractor().extract(
        "Chính sách thuê pin".encode(),
        filename="policy.txt",
        content_type="text/plain",
    )
    assert text == "Chính sách thuê pin"


@pytest.mark.asyncio
async def test_extracts_html_sections_tables_and_source_links_without_navigation_noise() -> None:
    text = await PolicyDocumentTextExtractor().extract(
        b"""
        <!doctype html><html><body>
          <nav>Trang chu San pham</nav>
          <h2>Bao hanh pin LFP</h2>
          <p>Ap dung cho hoa don tu ngay 16/08/2025.</p>
          <table><tr><th>Bo phan</th><th>Thoi han</th></tr>
          <tr><td>Pin LFP</td><td>8 nam</td></tr></table>
          <a href="/warranty.pdf">So bao hanh</a>
          <script>window.secret = 'ignore';</script>
        </body></html>
        """,
        filename="warranty-index.html",
        content_type="text/html",
    )

    assert "## Bao hanh pin LFP" in text
    assert "Bo phan | Thoi han" in text
    assert "Pin LFP | 8 nam" in text
    assert "So bao hanh (/warranty.pdf)" in text
    assert "Trang chu" not in text
    assert "window.secret" not in text


@pytest.mark.asyncio
async def test_extracts_docx_paragraphs_and_tables() -> None:
    stream = BytesIO()
    document = DocxDocument()
    document.add_paragraph("Battery warranty VF7")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Duration"
    table.cell(0, 1).text = "8 years"
    document.save(stream)

    text = await PolicyDocumentTextExtractor().extract(
        stream.getvalue(),
        filename="policy.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "Battery warranty VF7" in text
    assert "Duration | 8 years" in text


@pytest.mark.asyncio
async def test_extracts_text_based_pdf_without_ocr() -> None:
    stream = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=100)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}  # noqa: SLF001 - minimal PDF fixture
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 50 Td (Battery warranty VF7) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)  # noqa: SLF001 - minimal PDF fixture
    writer.write(stream)

    text = await PolicyDocumentTextExtractor().extract(
        stream.getvalue(),
        filename="policy.pdf",
        content_type="application/pdf",
    )

    assert "Battery warranty VF7" in text


@pytest.mark.asyncio
async def test_rejects_storage_supported_csv_for_policy_analysis() -> None:
    with pytest.raises(UnsupportedPolicyDocumentError):
        await PolicyDocumentTextExtractor().extract(
            b"name,value",
            filename="policy.csv",
            content_type="text/csv",
        )


@pytest.mark.asyncio
async def test_scanned_or_empty_pdf_fails_without_ocr() -> None:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(stream)

    with pytest.raises(DocumentTextUnavailableError):
        await PolicyDocumentTextExtractor().extract(
            stream.getvalue(),
            filename="scan.pdf",
            content_type="application/pdf",
        )
